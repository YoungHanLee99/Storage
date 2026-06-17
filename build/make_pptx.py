# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x0F, 0x2A, 0x4A)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
GRAY = RGBColor(0x55, 0x5F, 0x6B)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x20, 0x28, 0x32)
ACCENT = RGBColor(0xE8, 0x71, 0x22)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def box(slide, l, t, w, h, fill=None, line=None):
    sp = slide.shapes.add_shape(1, l, t, w, h)  # rectangle
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    return sp

def txt(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6, line_spacing=1.05):
    tb = slide.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05); tf.margin_top=Inches(0.02); tf.margin_bottom=Inches(0.02)
    first = True
    for item in runs:
        text, size, color, bold = item[0], item[1], item[2], (item[3] if len(item)>3 else False)
        bullet = item[4] if len(item)>4 else None
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        if bullet == 'i': p.level = 1
        r = p.add_run(); r.text = (("•  " if bullet=='b' else ("–  " if bullet=='i' else "")) + text)
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
        r.font.name = 'Malgun Gothic'
    return tb

def header(slide, kicker, title):
    box(slide, 0, 0, SW, Inches(1.15), fill=NAVY)
    box(slide, 0, Inches(1.15), SW, Pt(4), fill=ACCENT)
    txt(slide, Inches(0.6), Inches(0.16), Inches(12), Inches(0.35),
        [(kicker, 12, RGBColor(0x9F,0xC2,0xE0), True)])
    txt(slide, Inches(0.6), Inches(0.45), Inches(12.1), Inches(0.62),
        [(title, 25, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)

def footer(slide, n):
    txt(slide, Inches(0.6), Inches(7.05), Inches(8), Inches(0.3),
        [("d'strict  |  IPO Readiness Kick-off  |  2026-06-17  (대외비)", 9, GRAY)])
    txt(slide, Inches(12.0), Inches(7.05), Inches(0.9), Inches(0.3),
        [(str(n), 9, GRAY)], align=PP_ALIGN.RIGHT)

# ---------- Slide 1: Title ----------
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
box(s, 0, Inches(3.05), SW, Pt(4), fill=ACCENT)
txt(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.5),
    [("d'strict Holdings Inc.", 16, RGBColor(0x9F,0xC2,0xE0), True)])
txt(s, Inches(0.9), Inches(2.0), Inches(11.6), Inches(1.1),
    [("IPO Readiness Kick-off 회의 보고", 40, WHITE, True)])
txt(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(0.5),
    [("삼일PwC (Assurance / Global IPO) 착수 회의 결과 정리", 17, RGBColor(0xC7,0xD6,0xE6))])
txt(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(1.2),
    [("일시   2026-06-17(수) 09:30 ~ 10:05  ·  Microsoft Teams", 14, WHITE),
     ("참석   회사 3명 (이영한·장지만·김임영)  /  PwC 5명 (남승수·김승회·강찬휘·조범선·김혜진)", 13, RGBColor(0xC7,0xD6,0xE6)),
     ("작성   전략기획실(CSO/CLO)  ·  보안등급: 대외비", 12, GRAY)])

# ---------- Slide 2: 회의 개요 ----------
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
header(s, "MEETING OVERVIEW", "1.  회의 개요")
rows = [("회의명","d'strict_IPO Readiness (Kick-off)"),
        ("일시 / 방식","2026-06-17(수) 09:30 ~ 10:05 (35분)  ·  MS Teams"),
        ("주관","삼일PwC — Assurance / Global IPO"),
        ("목적","상장 옵션 검토 및 IPO Readiness 업무 범위 확정")]
y = Inches(1.55)
for k,v in rows:
    box(s, Inches(0.6), y, Inches(2.7), Inches(0.62), fill=NAVY)
    txt(s, Inches(0.72), y, Inches(2.5), Inches(0.62), [(k,13,WHITE,True)], anchor=MSO_ANCHOR.MIDDLE)
    box(s, Inches(3.3), y, Inches(9.4), Inches(0.62), fill=LIGHT)
    txt(s, Inches(3.5), y, Inches(9.1), Inches(0.62), [(v,13,DARK)], anchor=MSO_ANCHOR.MIDDLE)
    y = y + Inches(0.72)
# attendees two columns
box(s, Inches(0.6), Inches(4.65), Inches(6.0), Inches(2.0), fill=None, line=BLUE)
txt(s, Inches(0.8), Inches(4.78), Inches(5.6), Inches(0.4),[("■ 회사 (d'strict)",14,NAVY,True)])
txt(s, Inches(0.8), Inches(5.25), Inches(5.6), Inches(1.3),
    [("이영한 실장 — CSO/CLO (커뮤니케이션 창구)",12,DARK,False,'b'),
     ("장지만 실장 — CFO",12,DARK,False,'b'),
     ("김임영 팀장",12,DARK,False,'b')])
box(s, Inches(6.9), Inches(4.65), Inches(5.8), Inches(2.0), fill=None, line=BLUE)
txt(s, Inches(7.1), Inches(4.78), Inches(5.4), Inches(0.4),[("■ 자문사 (삼일PwC)",14,NAVY,True)])
txt(s, Inches(7.1), Inches(5.25), Inches(5.4), Inches(1.3),
    [("남승수 파트너 — Global IPO 총괄",12,DARK,False,'b'),
     ("김승회 · 강찬휘 · 조범선 (KR-ASR)",12,DARK,False,'b'),
     ("김혜진 — 감사 담당, Readiness 신규 합류",12,ACCENT,True,'b')])
footer(s, 2)

# ---------- Slide 3: 배경 및 목적 ----------
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
header(s, "BACKGROUND", "2.  회의 배경 및 목적")
txt(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(2.0),
    [("전통적으로 나스닥(미국) 상장을 우선 고려해 왔으나, 최근 한국 시장 상장도 옵션으로 진지하게 검토 중",15,DARK,False,'b'),
     ("본 회의는 상장 옵션 검토 및 IPO Readiness 업무 범위를 확정하기 위한 착수(Kick-off) 회의",15,DARK,False,'b'),
     ("피어그룹 분류 · 밸류에이션 스토리라인 · 감사/상장 요건을 중심으로 논의",15,DARK,False,'b')],
    space_after=12)
box(s, Inches(0.6), Inches(4.5), Inches(12.1), Inches(1.6), fill=LIGHT)
box(s, Inches(0.6), Inches(4.5), Pt(5), Inches(1.6), fill=ACCENT)
txt(s, Inches(0.9), Inches(4.7), Inches(11.6), Inches(1.3),
    [("계약 구조",12,BLUE,True),
     ("삼일PwC가 d'strict Holdings Inc.의 IPO 시장분석 자문사로 선정. 2025년 US-GAAP 재무제표 PA 업무와 IPO Readiness 업무를 통합한 단일 계약(통화 US$)으로 진행.",14,DARK)])
footer(s, 3)

# ---------- Slide 4: 주요 논의 (1) ----------
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
header(s, "KEY DISCUSSION (1/2)", "3.  주요 논의 내용 ①")
cards = [
 ("범위 · 산출물 · 일정", ["시장환경 분석 / 규제·요건 비교(美·韓)","피어 벤치마크 밸류에이션 시뮬레이션","상장 로드맵·마일스톤 제시","산출물: 美·韓 비교 PPT  ·  일정: 5주"]),
 ("상장 시장 선택 · 피어 분류", ["피어 분류가 밸류에이션에 핵심 영향","엔터테인먼트 vs 테크 하이브리드","한국이 美 대비 더 빠른 성장·상장 가능성"]),
]
x = Inches(0.6)
for title, items in cards:
    box(s, x, Inches(1.55), Inches(6.0), Inches(4.9), fill=None, line=BLUE)
    box(s, x, Inches(1.55), Inches(6.0), Inches(0.7), fill=NAVY)
    txt(s, x+Inches(0.2), Inches(1.55), Inches(5.6), Inches(0.7),[(title,15,WHITE,True)],anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+Inches(0.25), Inches(2.5), Inches(5.5), Inches(3.7),
        [(it,13.5,DARK,False,'b') for it in items], space_after=10, line_spacing=1.1)
    x = x + Inches(6.1)
footer(s, 4)

# ---------- Slide 5: 주요 논의 (2) ----------
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
header(s, "KEY DISCUSSION (2/2)", "3.  주요 논의 내용 ②")
cards = [
 ("사업계획 · 재무 현황", ["2030년까지 사업계획 보유","2026년 영업이익 흑자전환 목표","매출 성장 but 영업손실 지속","중간 CAPEX·대규모 파이낸싱 시점이 Pre-IPO 타이밍에 영향"]),
 ("상장 타이밍 · 성장 시점", ["PCAOB 업리프트 후 이론상 1년 내 US IPO 가능","단, 현실적 허들로 지연 가능성 큼","보고서에 '성장 시점'을 핵심 고려요소로 포함"]),
 ("두바이 법인 이슈", ["PCAOB 전환 시 예상 결함·개선포인트 추가 분석","두바이 매출 동향·결산 상황 점검"]),
 ("팀 · 커뮤니케이션", ["감사 담당 김혜진 합류","싱글 포인트(SPOC) 운영","회사 창구 공식화 + 김완 실장 CC"]),
]
positions = [(Inches(0.6),Inches(1.55)),(Inches(6.7),Inches(1.55)),(Inches(0.6),Inches(4.1)),(Inches(6.7),Inches(4.1))]
for (title,items),(x,y) in zip(cards,positions):
    box(s, x, y, Inches(6.0), Inches(2.4), fill=None, line=BLUE)
    box(s, x, y, Inches(6.0), Inches(0.55), fill=BLUE)
    txt(s, x+Inches(0.2), y, Inches(5.6), Inches(0.55),[(title,13.5,WHITE,True)],anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+Inches(0.25), y+Inches(0.62), Inches(5.5), Inches(1.7),
        [(it,11.5,DARK,False,'b') for it in items], space_after=5, line_spacing=1.05)
footer(s, 5)

# ---------- Slide 6: 핵심 쟁점 & 결정사항 ----------
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
header(s, "ISSUES & DECISIONS", "4.  핵심 쟁점 및 결정 사항")
box(s, Inches(0.6), Inches(1.55), Inches(6.0), Inches(4.9), fill=None, line=GRAY)
box(s, Inches(0.6), Inches(1.55), Inches(6.0), Inches(0.6), fill=NAVY)
txt(s, Inches(0.8), Inches(1.55), Inches(5.6), Inches(0.6),[("핵심 쟁점 (Key Questions)",14,WHITE,True)],anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.85), Inches(2.35), Inches(5.5), Inches(4.0),
    [("미국(나스닥) vs 한국 — 우선 시장은?",12.5,DARK,False,'b'),
     ("지금이 착수에 적절한 시점인가?",12.5,DARK,False,'b'),
     ("성장 목표 시점과 시장 선택 접근법",12.5,DARK,False,'b'),
     ("PCAOB 업리프트 후 1년 내 US IPO 가능?",12.5,DARK,False,'b'),
     ("두바이 PCAOB 업리프트 소요기간은?",12.5,DARK,False,'b'),
     ("두바이 매출·결산 현황은?",12.5,DARK,False,'b')], space_after=11)
box(s, Inches(6.9), Inches(1.55), Inches(5.8), Inches(4.9), fill=LIGHT)
box(s, Inches(6.9), Inches(1.55), Inches(5.8), Inches(0.6), fill=ACCENT)
txt(s, Inches(7.1), Inches(1.55), Inches(5.4), Inches(0.6),[("결정 사항",14,WHITE,True)],anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(7.15), Inches(2.35), Inches(5.4), Inches(4.0),
    [("업무 범위 4개 영역 · 5주 일정 확정",12.5,DARK,True,'b'),
     ("(시장분석·요건비교·밸류에이션·로드맵)",11,GRAY,False,'i'),
     ("보고서에 '성장 시점(Timing)' 반영",12.5,DARK,True,'b'),
     ("두바이 법인 결함·개선포인트 분석 포함",12.5,DARK,True,'b'),
     ("커뮤니케이션 싱글포인트 운영 (김완 실장 CC)",12.5,DARK,True,'b')], space_after=12)
footer(s, 6)

# ---------- Slide 7: Action Items ----------
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
header(s, "ACTION ITEMS", "5.  후속 조치 (Action Items)")
# company table
txt(s, Inches(0.6), Inches(1.4), Inches(6), Inches(0.4),[("■ 회사 (d'strict)",14,NAVY,True)])
comp = [("장지만 CFO","2030 사업계획(시나리오) 추가 전달"),
        ("장지만 CFO","재무추정·자금조달 자료 전달 (회의 직후)"),
        ("이영한 실장","업데이트된 회의 자료 준비·공유"),
        ("이영한 실장","IR 자료·2030 사업계획서·자회사 정보 공유"),
        ("이영한 실장","인수 대상/계획 정보 공유 (있을 경우)"),
        ("이영한 실장","SPOC 연락처 공식화 + 김완 실장 CC")]
y = Inches(1.85)
for who,what in comp:
    box(s, Inches(0.6), y, Inches(2.3), Inches(0.55), fill=LIGHT)
    txt(s, Inches(0.7), y, Inches(2.2), Inches(0.55),[(who,11,NAVY,True)],anchor=MSO_ANCHOR.MIDDLE)
    box(s, Inches(2.9), y, Inches(3.7), Inches(0.55), fill=None, line=RGBColor(0xD5,0xDD,0xE5))
    txt(s, Inches(3.0), y, Inches(3.5), Inches(0.55),[(what,10.5,DARK)],anchor=MSO_ANCHOR.MIDDLE)
    y = y + Inches(0.62)
# pwc table
txt(s, Inches(7.0), Inches(1.4), Inches(6), Inches(0.4),[("■ 자문사 (삼일PwC)",14,ACCENT,True)])
pwc = [("남승수","PBC(요청자료) 리스트 별도 송부"),
       ("남승수","보고서에 '성장 시점' 주요 고려요소 포함"),
       ("남승수","두바이 법인 결함·개선포인트 분석 수행"),
       ("남승수","진행 중 추가 질문·요청 지속 전달")]
y = Inches(1.85)
for who,what in pwc:
    box(s, Inches(7.0), y, Inches(1.7), Inches(0.55), fill=RGBColor(0xFB,0xE9,0xDA))
    txt(s, Inches(7.1), y, Inches(1.6), Inches(0.55),[(who,11,ACCENT,True)],anchor=MSO_ANCHOR.MIDDLE)
    box(s, Inches(8.7), y, Inches(4.0), Inches(0.55), fill=None, line=RGBColor(0xD5,0xDD,0xE5))
    txt(s, Inches(8.8), y, Inches(3.8), Inches(0.55),[(what,10.5,DARK)],anchor=MSO_ANCHOR.MIDDLE)
    y = y + Inches(0.62)
footer(s, 7)

# ---------- Slide 8: 리스크 & 참고자료 ----------
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
header(s, "RISKS & REFERENCES", "6.  리스크 및 참고자료")
box(s, Inches(0.6), Inches(1.55), Inches(6.0), Inches(4.9), fill=None, line=ACCENT)
box(s, Inches(0.6), Inches(1.55), Inches(6.0), Inches(0.6), fill=ACCENT)
txt(s, Inches(0.8), Inches(1.55), Inches(5.6), Inches(0.6),[("리스크 · 점검 포인트",14,WHITE,True)],anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.85), Inches(2.4), Inches(5.5), Inches(4.0),
    [("영업손실 지속 → 트랙별 재무·감사 요건 충족이 관건",12.5,DARK,False,'b'),
     ("피어 분류(엔터 vs 테크)에 따라 밸류 편차 큼",12.5,DARK,False,'b'),
     ("두바이 법인 PCAOB 전환 결함·업리프트 기간 불확실",12.5,DARK,False,'b'),
     ("'1년 내 IPO'는 이론치 — 현실적 지연 요인 다수",12.5,DARK,False,'b')], space_after=14)
box(s, Inches(6.9), Inches(1.55), Inches(5.8), Inches(4.9), fill=LIGHT)
box(s, Inches(6.9), Inches(1.55), Inches(5.8), Inches(0.6), fill=NAVY)
txt(s, Inches(7.1), Inches(1.55), Inches(5.4), Inches(0.6),[("참고 자료 (Google Drive)",14,WHITE,True)],anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(7.15), Inches(2.4), Inches(5.4), Inches(4.0),
    [("d'strict IPO Readiness Proposal_PwC.pdf",11.5,DARK,True,'b'),
     ("260202_나스닥 상장 IPO",11.5,DARK,False,'b'),
     ("나스닥 상장 유니콘 IPO 준비 로드맵",11.5,DARK,False,'b'),
     ("디스트릭트 IPO 준비 및 신사업 전략(비전 2028)",11.5,DARK,False,'b'),
     ("241017_IPO 자금조달 시뮬레이션.xlsx",11.5,DARK,False,'b'),
     ("240930_대주주 IPO품의건_v3.xlsx",11.5,DARK,False,'b')], space_after=11)
footer(s, 8)

prs.save('build/d_strict_IPO_Readiness_Kickoff_20260617.pptx')
print("saved pptx")
